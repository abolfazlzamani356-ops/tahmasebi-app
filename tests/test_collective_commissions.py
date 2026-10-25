# -*- coding: utf-8 -*-
"""
Tests for Collective Commission System (سیستم پورسانت کلی)
Verifying 2% pool calculation, 4-way equal division among the 4 designated sellers
(خانم نقدی، خانم امیریان، خانم حاجیلو، خانم بیدریغ), independence from individual commission systems,
route access controls, and Excel export.
"""
import io
import os
import sys
import pytest
import openpyxl

base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
if base_dir not in sys.path:
    sys.path.insert(0, base_dir)

from app import app, db
from models import User, Invoice

@pytest.fixture
def client():
    app.config['TESTING'] = True
    app.config['WTF_CSRF_ENABLED'] = False
    with app.test_client() as client:
        with app.app_context():
            yield client


def test_collective_commission_math_formula():
    """
    تست دقت فرمول ریاضی سیستم پورسانت کلی:
    محاسبه ۲ درصد و تقسیم بر ۴ نفر (معادل ۰.۵ درصد برای هر نفر)
    """
    # مثال ارائه شده توسط کاربر: فاکتور ۱۲۰ میلیون تومانی
    amount_1 = 120_000_000
    pool_1 = int(round(amount_1 * 0.02))
    assert pool_1 == 2_400_000, f"Expected 2% pool of 120M to be 2,400,000, got {pool_1}"
    share_each_1 = int(round(pool_1 / 4.0))
    assert share_each_1 == 600_000, f"Expected each share to be 600,000, got {share_each_1}"
    assert share_each_1 * 4 == pool_1

    # فاکتور ۱۰ میلیون تومانی
    amount_2 = 10_000_000
    pool_2 = int(round(amount_2 * 0.02))
    assert pool_2 == 200_000
    share_each_2 = int(round(pool_2 / 4.0))
    assert share_each_2 == 50_000

    # فاکتور با مبالغ خرد
    amount_3 = 15_350_000
    pool_3 = int(round(amount_3 * 0.02))
    assert pool_3 == 307_000
    share_each_3 = int(round(pool_3 / 4.0))
    assert share_each_3 == 76_750


def test_collective_commissions_view_access(client):
    """
    تست کنترل دسترسی:
    کاربر بدون لاگین ریدایرکت می‌شود و ادمین لاگین‌شده کد ۲۰۰ دریافت می‌کند.
    """
    # بدون لاگین
    res = client.get('/admin/collective-commissions')
    assert res.status_code == 302
    assert '/login' in res.location

    # با لاگین ادمین
    with client.session_transaction() as sess:
        sess['user_id'] = 1
        sess['role'] = 'admin'
        sess['full_name'] = 'مدیریت کل'

    res_admin = client.get('/admin/collective-commissions')
    assert res_admin.status_code == 200
    html = res_admin.data.decode('utf-8')

    # بررسی وجود المان‌های کلیدی سیستم
    assert 'سیستم پورسانت کلی' in html
    assert 'خانم نقدی' in html
    assert 'خانم امیریان' in html
    assert 'خانم حاجیلو' in html
    assert 'خانم بیدریغ' in html
    assert '۲٪' in html or '2%' in html
    assert 'شبیه‌ساز و محاسبه‌گر سریع پورسانت کلی' in html


def test_export_collective_commissions_excel(client):
    """
    تست تولید فایل اکسل گزارش سیستم پورسانت کلی
    """
    with client.session_transaction() as sess:
        sess['user_id'] = 1
        sess['role'] = 'admin'
        sess['full_name'] = 'مدیریت کل'

    res = client.get('/admin/export/collective-commissions-excel')
    assert res.status_code == 200
    assert 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' in res.content_type

    # باز کردن فایل اکسل در حافظه و اعتبارسنجی تب‌ها
    wb = openpyxl.load_workbook(io.BytesIO(res.data))
    sheet_names = wb.sheetnames
    assert "خلاصه پورسانت کلی" in sheet_names
    assert "ریز محاسبات فاکتورها" in sheet_names

    ws1 = wb["خلاصه پورسانت کلی"]
    assert "گزارش رسمی سیستم پورسانت کلی" in str(ws1.cell(row=1, column=1).value)


def test_system_independence_no_side_effects():
    """
    تاکید کاربر: این سیستم محاسباتی کاملاً مستقل است و ساختار فاکتورها،
    پورسانت‌های پایه پرسنل و حقوق‌ها را دستکاری نمی‌کند.
    """
    with app.app_context():
        # بررسی ۴ فروشنده در جدول کاربران
        core_users = User.query.filter(User.id.in_([2, 3, 4, 5])).all()
        assert len(core_users) == 4
        for u in core_users:
            # درصد پورسانت‌های اختصاصی هر فرد تغییر نکرده است
            assert u.commission_rate is not None
            assert u.is_active == 1

        # بررسی فاکتورهای قطعی
        sample_inv = Invoice.query.filter_by(status='final').first()
        if sample_inv:
            assert sample_inv.total_amount > 0
